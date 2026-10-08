#Bibliotecas
import pandas as pd
import numpy as np
from datetime import date
from openpyxl import load_workbook
import streamlit as st
import string

st.title("Control de Pedidos", text_alignment="center")

st.divider()
st.subheader("Seleccionar el numero de bases activas")
numero_bases_activas = st.number_input("Numero de Bases activas?", min_value=1, max_value=10)

st.divider()
st.subheader("Seleccionar el numero columnas requeridas para pedidos")
numero_de_pedidos = st.number_input("Numero de columnas de Pedidos?", min_value=1, max_value=5)

st.divider()
st.header("Archvios necesarios para el procesamiento:")
lista_de_materiales = st.file_uploader("Seleccionar archivo con la LISTA DE MATERIALES necesarios en el archivos de pedidos:", type=["xlsx","xls"])
inventario_sae_ = st.file_uploader("Seleccionar INVENTARIO DE SAE:", type=["xlsx","xls"])
inventario_homologado = st.file_uploader("Seleccionar INVENTARIO HOMOLOGADO SEMANAL:", type=["xlsx","xls", "xlsm"])
st.divider()


if lista_de_materiales is not None and inventario_sae_ is not None and inventario_homologado is not None:
  # Leer el archivo Excel de Lista de Materiales e Inventario SAE
  df = pd.read_excel(lista_de_materiales)
  inventario_sae = pd.read_excel(inventario_sae_)

  st.write("### Lista de Materiales")
  st.dataframe(df)

  st.write("### Inventario SAE")
  st.dataframe(inventario_sae)

  # Ingenieria de variables para posterior lectura del archivo de Inventario homologado semanal
  nombre_bases_activas = pd.ExcelFile(inventario_homologado).sheet_names[1:numero_bases_activas+1]
  columnas = list(np.array(pd.read_excel(inventario_homologado)[3:4])[0][:2]) + list(np.array(pd.read_excel(inventario_homologado)[2:3])[0][2:])
  #Agregamos la base a las columnas especificas de cada base
  col=1
  for i in nombre_bases_activas:
      for j in range(1,5):
          columnas[col+j] = columnas[col+j] + f' {i}'
      col=j+1

  #Leemos el archivo de Inventario homologado semanal
  homologado_inventarios_bases = pd.read_excel(inventario_homologado)[4:].dropna(subset=['Unnamed: 0'])
  homologado_inventarios_bases.columns = columnas
  homologado_inventarios_bases.rename(columns={'Total\n+\nCotizaciones':'Total en Bases'}, inplace=True)
  homologado_inventarios_bases.reset_index(drop=True, inplace=True)
  
  st.write("### Inventario Homologado semanal")
  st.dataframe(homologado_inventarios_bases)

  #Procesamiento de la tabla para el archivo principal
  if st.button("Crear archivo de Pedidos"):

    diccionario_letras = {i:letra for i, letra in enumerate(string.ascii_uppercase, start=1)}
  
    #Creamos un diccionario con las tablas de los inventarios de cada base activa
    tablas_bases = {}
    c=2
    for i in nombre_bases_activas:
        tablas_bases[i]=homologado_inventarios_bases[['NO. DE PARTE ', homologado_inventarios_bases.columns[c], homologado_inventarios_bases.columns[c+1], homologado_inventarios_bases.columns[c+2], homologado_inventarios_bases.columns[c+3]]]
        c+=4

    #Creamos diccionarios con la informacion de SAE de los materiales
    dictio_descripcion = dict(zip(inventario_sae['Clave '],inventario_sae['Descripción ']))
    dictio_linea = dict(zip(inventario_sae['Clave '],inventario_sae['Línea ']))
    dictio_inventario_sae = dict(zip(inventario_sae['Clave '],inventario_sae['Existencias ']))

    #Agregamos a la tabla principal la descripcion y la existencia en SAE
    df['Descripcion'] = df['Numero de Parte'].map(lambda x: dictio_descripcion[str(x)])
    df[f'Existencia SAE {date.today()}'] = df['Numero de Parte'].map(lambda x: dictio_inventario_sae[str(x)])

    #Agregamos la informacion de las bases activas a la tabla principal
    for i in nombre_bases_activas:
      df = df.merge(tablas_bases[i], how='left', left_on='Numero de Parte', right_on='NO. DE PARTE ').drop(['NO. DE PARTE '], axis=1).copy()
      
    #Creamos un diccionario para numerar las columnas
    diccionario_columnas = {columnas:i for i, columnas in enumerate(df.columns, start=1)}

    #Creamos una lista con los nombres de las columnas de suma de cada base
    c=6
    cols_sumas_bases=[]
    for i in range(numero_bases_activas):
      cols_sumas_bases.append(df.columns[c])
      c+=4

    #Agregamos columna de TOTAL de materiales en las bases y sumamos las existencias de las bases activas
    formula_suma_bases = "= +"
    for i in cols_sumas_bases:
      formula_suma_bases += diccionario_letras[diccionario_columnas[i]] + '{} +'
    formula_suma_bases = formula_suma_bases[:-2]
    df['TOTAL en Bases'] = [formula_suma_bases.format(j, j, j, j, j, j, j, j, j, j) for j in range(2, df.shape[0]+2)]

    st.write(df['TOTAL en Bases'])

    c=6
    for i in nombre_bases_activas:
        df[df.columns[c]] = [0 if str(x)=='nan' else x for x in df[df.columns[c]]]
        c+=4

    c=6
    for i in nombre_bases_activas:
        df['TOTAL en Bases'] = df['TOTAL en Bases'] + df[df.columns[c]]
        c+=4

    #Agregamos la variable de TOTAL de existencias de SAE menos las existencias en las Bases activas
    df['TOTAL: SAE - Bases'] = df[f'Existencia SAE {date.today()}'] - df['TOTAL en Bases']

    # CAMBIAR A DINAMICO Agregamos la columna de cada pedido as como el total de los pedidos CAMBIAR A DINAMICO
    for i in range(1, numero_de_pedidos + 1):
      df[f'PEDIDO {i}:'] = [0 for x in range(df.shape[0])]

    diccionario_columnas = {columnas:i for i, columnas in enumerate(df.columns, start=1)}

    formula_suma_pedidos = "= +"
    for i in range(1, numero_de_pedidos + 1):
      formula_suma_pedidos += diccionario_letras[diccionario_columnas[f'PEDIDO {i}:']] + '{} +'
    formula_suma_pedidos = formula_suma_pedidos[:-2]

    df['TOTAL Pedidos'] = [formula_suma_pedidos.format(j, j, j, j, j) for j in range(2, df.shape[0]+2)]
    
    diccionario_columnas = {columnas:i for i, columnas in enumerate(df.columns, start=1)}

    #Agregamos la columna del TOTAL: (Pedidos) + (SAE - Bases)
    df['TOTAL: (Pedidos) + (SAE-Bases)'] = ['= + ' + diccionario_letras[diccionario_columnas['TOTAL Pedidos']] + f'{j} + ' + diccionario_letras[diccionario_columnas['TOTAL: SAE - Bases']] + f'{j}' for j in range(2, df.shape[0]+2)]
    
    #Reemplazamos los ceros por nulos para que no hagan ruido en el archivo
    #df.replace(0, np.nan, inplace=True)

    st.dataframe(df)

    # Permitir descargar el resultado
    output_name = f"Pedidos_{date.today()}.xlsx"
    df.to_excel(output_name, index=False)

    #Editamos el archivo para ajustar las celdas al texto
    wb = load_workbook(output_name)
    ws = wb[wb.sheetnames[0]]
    for col in ws.columns:
      max_length = 0
      col_letter = col[0].column_letter
      for cell in col:
        try:
          if cell.value:
            max_length = max(max_length, len(str(cell.value)))
        except:
          pass
      # Asignar un margen adicional para que no quede apretado
      ws.column_dimensions[col_letter].width = max(max_length + 3, 10)
    wb.save(output_name)
    
    with open(output_name, "rb") as file:
        st.download_button(
            label="Descargar Archivo de pedidos",
            data=file,
            file_name=f"Pedidos_{date.today()}.xlsx",
            mime=f"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )


    










